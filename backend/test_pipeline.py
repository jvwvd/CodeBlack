import unittest
from io import BytesIO
from unittest.mock import patch

import pymupdf
from docx import Document
from openpyxl import Workbook

from pipeline import reliability
from pipeline.classify import classify_email
from pipeline.documents import (
    identify_doc_type,
    parse_attachment_bytes,
)
from pipeline.run import _finalize, process_email
from models import EmailResult


SI_TEXT = """
SHIPPING INSTRUCTION

Shipper: TEST EXPORTER LTD
Consignee: TEST IMPORTER SDN BHD
Notify Party: TEST IMPORTER SDN BHD
Port of Loading: SINGAPORE
Port of Discharge: KARACHI, PAKISTAN (PKKHI)
Container Count: 3
Gross Weight: 22000 KG
"""

BL_TEXT = """
DRAFT BILL OF LADING

Shipper: TEST EXPORTER LTD
Consignee: TEST IMPORTER SDN BHD
Notify Party: TEST IMPORTER SDN BHD
Port of Loading: SINGAPORE
Port of Discharge: KARACHI, PAKISTAN (PKKHI)
Container Count: 3
Gross Weight: 22000 KG
"""


def make_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def make_empty_pdf() -> bytes:
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(text: str) -> bytes:
    document = Document()

    for line in text.strip().splitlines():
        document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(text: str) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active

    for row_index, line in enumerate(text.strip().splitlines(), start=1):
        worksheet.cell(row=row_index, column=1, value=line)

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    return buffer.getvalue()


class PipelineRegressionTests(unittest.TestCase):

    def test_txt_si_bl_end_to_end(self):
        email = {
            "email_id": "email_test_txt",
            "from": "ops@example.com",
            "subject": "Please check draft BL against shipping instruction",
            "body": "Please verify the shipping instruction against the draft bill of lading.",
            "attachments": [
                "attachments/si.txt",
                "attachments/bl.txt",
            ],
        }

        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
            "attachments/bl.txt": BL_TEXT.encode(),
        }

        result = process_email(
            email,
            attachment_bytes=attachment_bytes,
        )

        self.assertEqual(result.category, "BL_COMPARISON")
        self.assertEqual(result.status, "OK")
        self.assertFalse(result.has_defect)
        self.assertEqual(result.defect_fields, [])

    def test_pdf_text_extraction(self):
        data = make_pdf(BL_TEXT)

        text = parse_attachment_bytes(
            data,
            "draft_bl.pdf",
        )

        self.assertIsNotNone(text)
        self.assertIn("DRAFT BILL OF LADING", text)
        self.assertEqual(identify_doc_type(text), "BL")

    def test_docx_extraction(self):
        data = make_docx(BL_TEXT)

        text = parse_attachment_bytes(
            data,
            "draft_bl.docx",
        )

        self.assertIsNotNone(text)
        self.assertIn("DRAFT BILL OF LADING", text)
        self.assertEqual(identify_doc_type(text), "BL")

    def test_xlsx_extraction(self):
        data = make_xlsx(SI_TEXT)

        text = parse_attachment_bytes(
            data,
            "shipping_instruction.xlsx",
        )

        self.assertIsNotNone(text)
        self.assertIn("SHIPPING INSTRUCTION", text)
        self.assertEqual(identify_doc_type(text), "SI")

    @patch(
        "pipeline.documents.read_scanned_pdf_with_vision",
        return_value=SI_TEXT,
    )
    def test_scanned_pdf_uses_vision_fallback(self, mock_vision):
        data = make_empty_pdf()

        text = parse_attachment_bytes(
            data,
            "scanned_si.pdf",
        )

        mock_vision.assert_called_once()

        self.assertIsNotNone(text)
        self.assertIn("SHIPPING INSTRUCTION", text)
        self.assertEqual(identify_doc_type(text), "SI")

    def test_missing_attachment_routes_to_review(self):
        email = {
            "email_id": "email_test_missing",
            "from": "ops@example.com",
            "subject": "Please check draft BL against shipping instruction",
            "body": "Please verify the shipping instruction against the draft bill of lading.",
            "attachments": [
                "attachments/si.txt",
            ],
        }

        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
        }

        category, _ = classify_email(email)

        self.assertEqual(
            category,
            "BL_COMPARISON",
        )

        result = process_email(
            email,
            attachment_bytes=attachment_bytes,
        )

        self.assertEqual(
            result.status,
            "NEEDS_REVIEW",
        )

        self.assertEqual(
            result.review_reason,
            "missing_attachment",
        )

    @patch(
        "pipeline.classify.classify_with_llm",
        return_value=None,
    )
    def test_ai_unavailable_has_safe_classification_fallback(
        self,
        mock_llm,
    ):
        email = {
            "email_id": "email_test_ai_failure",
            "from": "ops@example.com",
            "subject": "Operational update",
            "body": "Please review when available.",
            "attachments": [],
        }

        category, decided_by = classify_email(email)

        mock_llm.assert_called_once()

        self.assertEqual(category, "GENERAL")
        self.assertEqual(decided_by, "rule")

    @patch(
        "pipeline.documents.read_scanned_pdf_with_vision",
        return_value=None,
    )
    def test_failed_vision_returns_unreadable(
        self,
        mock_vision,
    ):
        data = make_empty_pdf()

        text = parse_attachment_bytes(
            data,
            "unreadable_scan.pdf",
        )

        mock_vision.assert_called_once()
        self.assertIsNone(text)

    def test_si_request_subject_beats_invoice_text_in_body(self):
        body = """
        Please find the shipping instruction attached.

        Documents Required:
        1) 3 Original invoice
        2) Packing list
        3) Original BL

        Forwarded thread also mentions billing and payment.
        """

        subjects = (
            "CUST SI _ MEA _ 5RCY-52735 __ PO_25_5465",
            "RE_ CUST SI _ MEA _ 5RCY-51168 __ PO_25_3508",
            "REQUEST SI _ 5APH-62718 _ HOUSTON_US",
            "RE_ REQUEST SI _ 5RCY-60883 _ CONAKRY_GUINEA",
            "SI NEEDED_ 5RUS-16202 _ EAST BRIGHT",
            "RE_ SI NEEDED_ 5APH-26773 _ UAB NOVAKOPA",
            "SI - HLCUSIN331541006 - DIRECT(HAPAG) - 5RUS-61793",
            "RE_ SI - SIJ3777014 - DIRECT(CMA) - 5ALT-88568",
        )

        for subject in subjects:
            with self.subTest(subject=subject):
                email = {
                    "email_id": "email_test_si_request",
                    "from": "ops@example.com",
                    "subject": subject,
                    "body": body,
                    "attachments": [],
                }

                category, decided_by = classify_email(email)

                self.assertEqual(category, "SI_REQUEST")
                self.assertEqual(decided_by, "rule")

    def test_documented_invoice_subject_signals(self):
        subjects = (
            "REQUEST TO CANCEL INVOICE -5250075802",
            "RE_ LOCAL CHARGES FOB - KARGOSMAR - TELEX RELEASE CHARGES",
            "2157 RAK BILLING 5070146693 MISSING GR",
            "_RPA_ India HSS SD Billing Process Completed - VISION",
            "Mill D & D charges - 6437419230",
        )

        for subject in subjects:
            with self.subTest(subject=subject):
                email = {
                    "email_id": "email_test_invoice",
                    "from": "ops@example.com",
                    "subject": subject,
                    "body": "Normal forwarded operational thread.",
                    "attachments": [],
                }

                category, decided_by = classify_email(email)

                self.assertEqual(category, "INVOICE_QUERY")
                self.assertEqual(decided_by, "rule")


SI_TEXT_MISSING_WEIGHT = """
SHIPPING INSTRUCTION

Shipper: TEST EXPORTER LTD
Consignee: TEST IMPORTER SDN BHD
Notify Party: TEST IMPORTER SDN BHD
Port of Loading: SINGAPORE
Port of Discharge: KARACHI, PAKISTAN (PKKHI)
Container Count: 3
"""

BL_TEXT_DIFFERENT_CONSIGNEE = """
DRAFT BILL OF LADING

Shipper: TEST EXPORTER LTD
Consignee: DIFFERENT IMPORTER SDN BHD
Notify Party: TEST IMPORTER SDN BHD
Port of Loading: SINGAPORE
Port of Discharge: KARACHI, PAKISTAN (PKKHI)
Container Count: 3
Gross Weight: 22000 KG
"""


class DefectReviewInvariantTests(unittest.TestCase):
    """Organizer README invariant: has_defect/defect_fields are reported
    for MISMATCH, review_reason is reported for NEEDS_REVIEW -- never
    both. Regression coverage for email_040/email_516/email_518, which
    previously came back NEEDS_REVIEW + review_reason=missing_value while
    still carrying has_defect=true and a non-empty defect_fields list."""

    @patch("pipeline.extract.extract_with_llm", return_value=None)
    def test_missing_si_weight_with_apparent_differences_is_clean_needs_review(self, mock_llm):
        # SI is missing gross weight (required field) AND consignee genuinely
        # differs from BL -- the exact shape that used to leak has_defect/
        # defect_fields alongside NEEDS_REVIEW.
        email = {
            "email_id": "email_test_missing_weight",
            "from": "ops@example.com",
            "subject": "Please check draft BL against shipping instruction",
            "body": "Please verify the shipping instruction against the draft bill of lading.",
            "attachments": ["attachments/si.txt", "attachments/bl.txt"],
        }
        attachment_bytes = {
            "attachments/si.txt": SI_TEXT_MISSING_WEIGHT.encode(),
            "attachments/bl.txt": BL_TEXT_DIFFERENT_CONSIGNEE.encode(),
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertEqual(result.status, "NEEDS_REVIEW")
        self.assertEqual(result.review_reason, "missing_value")
        self.assertFalse(result.has_defect)
        self.assertEqual(result.defect_fields, [])

    def test_real_mismatch_with_complete_data_still_reports_defects(self):
        # Both sides fully readable/complete, genuine difference on consignee:
        # must remain MISMATCH with has_defect/defect_fields populated and
        # no review_reason.
        email = {
            "email_id": "email_test_real_mismatch",
            "from": "ops@example.com",
            "subject": "Please check draft BL against shipping instruction",
            "body": "Please verify the shipping instruction against the draft bill of lading.",
            "attachments": ["attachments/si.txt", "attachments/bl.txt"],
        }
        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
            "attachments/bl.txt": BL_TEXT_DIFFERENT_CONSIGNEE.encode(),
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertEqual(result.status, "MISMATCH")
        self.assertIsNone(result.review_reason)
        self.assertTrue(result.has_defect)
        self.assertIn("consignee", result.defect_fields)

    def test_exact_match_remains_ok_with_no_defect_markers(self):
        email = {
            "email_id": "email_test_exact_match",
            "from": "ops@example.com",
            "subject": "Please check draft BL against shipping instruction",
            "body": "Please verify the shipping instruction against the draft bill of lading.",
            "attachments": ["attachments/si.txt", "attachments/bl.txt"],
        }
        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
            "attachments/bl.txt": BL_TEXT.encode(),
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertEqual(result.status, "OK")
        self.assertIsNone(result.review_reason)
        self.assertFalse(result.has_defect)
        self.assertEqual(result.defect_fields, [])

    def test_finalize_strips_defects_whenever_review_reason_is_set(self):
        # Direct unit coverage of the guard itself, independent of how the
        # (possibly stale) has_defect/defect_fields got onto the result.
        leaky = EmailResult(
            email_id="email_leaky",
            category="BL_COMPARISON",
            status="NEEDS_REVIEW",
            defect_fields=["consignee", "gross_weight_kg"],
            has_defect=True,
            review_reason="missing_value",
        )

        cleaned = _finalize(leaky)

        self.assertEqual(cleaned.status, "NEEDS_REVIEW")
        self.assertFalse(cleaned.has_defect)
        self.assertEqual(cleaned.defect_fields, [])

    def test_finalize_leaves_mismatch_result_untouched(self):
        clean = EmailResult(
            email_id="email_clean",
            category="BL_COMPARISON",
            status="MISMATCH",
            defect_fields=["shipper"],
            has_defect=True,
            review_reason=None,
        )

        result = _finalize(clean)

        self.assertEqual(result.status, "MISMATCH")
        self.assertTrue(result.has_defect)
        self.assertEqual(result.defect_fields, ["shipper"])


WRONG_DOC_TEXT = """
COMMERCIAL INVOICE

Invoice No.: INV-2026-001
Seller: TEST EXPORTER LTD
Buyer: TEST IMPORTER SDN BHD

*** THIS IS A COMMERCIAL INVOICE - NOT A SHIPPING INSTRUCTION ***
"""

COMPARISON_SUBJECT = "TO CONFIRM DOCS _ Please check draft BL against shipping instruction"
COMPARISON_BODY = "Please verify the shipping instruction against the draft bill of lading."

SEND_DRAFT_BL_SUBJECT = "REQUEST BL DRAFT _ PO 25041 _ PAPERONE DIGITAL COPIER PAPER"
SEND_DRAFT_BL_BODY = "Dear Team,\n\nPlease assist to send the draft BL for booking 070500208599 for checking asap.\n\nThank you."

COMPARE_NO_ATTACHMENT_BODY = (
    "Dear Team,\n\nPlease compare the SI and draft BL for 070500263211 and confirm "
    "(attachments appear to have been dropped). Thank you."
)


class MissingAttachmentReliabilityTests(unittest.TestCase):
    """Regression coverage for the missing_attachment over-triggering bug:
    reliability.check() used to treat ANY BL_COMPARISON email with fewer
    than 2 attachment references as a genuinely missing document, which
    also flagged ordinary "please send the draft BL" requests (which never
    had attachments to begin with -- see the organizer README: emails with
    no attachments are a normal, expected case, not a review case)."""

    # -- direct unit coverage of reliability.check() -----------------------

    def test_zero_attachments_without_si_reference_is_not_missing_attachment(self):
        email = {
            "email_id": "email_send_bl",
            "subject": SEND_DRAFT_BL_SUBJECT,
            "body": SEND_DRAFT_BL_BODY,
            "attachments": [],
        }
        self.assertIsNone(reliability.check(email, None, None))

    def test_zero_attachments_with_si_reference_is_missing_attachment(self):
        email = {
            "email_id": "email_dropped",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARE_NO_ATTACHMENT_BODY,
            "attachments": [],
        }
        self.assertEqual(reliability.check(email, None, None), "missing_attachment")

    def test_one_attachment_is_always_missing_attachment(self):
        email = {
            "email_id": "email_partial",
            "subject": SEND_DRAFT_BL_SUBJECT,
            "body": SEND_DRAFT_BL_BODY,  # no SI reference at all
            "attachments": ["attachments/si.txt"],
        }
        # Partial (exactly one of two required documents) is always genuine,
        # regardless of whether the email text happens to mention "SI".
        self.assertEqual(reliability.check(email, SI_TEXT, None), "missing_attachment")

    # -- end-to-end process_email() coverage --------------------------------

    def test_1_genuine_missing_attachment_only_one_physical_document(self):
        email = {
            "email_id": "email_test_one_doc",
            "from": "ops@example.com",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARISON_BODY,
            "attachments": ["attachments/si.txt"],
        }
        attachment_bytes = {"attachments/si.txt": SI_TEXT.encode()}

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertEqual(result.status, "NEEDS_REVIEW")
        self.assertEqual(result.review_reason, "missing_attachment")
        self.assertFalse(result.has_defect)
        self.assertEqual(result.defect_fields, [])

    def test_2_weak_filenames_with_identifiable_content_not_missing_attachment(self):
        email = {
            "email_id": "email_test_weak_filenames",
            "from": "ops@example.com",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARISON_BODY,
            "attachments": ["attachments/doc1.txt", "attachments/doc2.txt"],
        }
        attachment_bytes = {
            "attachments/doc1.txt": SI_TEXT.encode(),
            "attachments/doc2.txt": BL_TEXT.encode(),
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertNotEqual(result.review_reason, "missing_attachment")
        self.assertEqual(result.status, "OK")
        self.assertIsNone(result.review_reason)

    def test_3_two_files_wrong_doc_type_not_missing_attachment(self):
        email = {
            "email_id": "email_test_wrong_doc",
            "from": "ops@example.com",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARISON_BODY,
            "attachments": ["attachments/si.txt", "attachments/invoice.txt"],
        }
        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
            "attachments/invoice.txt": WRONG_DOC_TEXT.encode(),
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertEqual(result.review_reason, "wrong_doc_type")
        self.assertNotEqual(result.review_reason, "missing_attachment")
        self.assertEqual(result.status, "NEEDS_REVIEW")

    def test_4_unreadable_file_not_missing_attachment(self):
        email = {
            "email_id": "email_test_unreadable",
            "from": "ops@example.com",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARISON_BODY,
            "attachments": ["attachments/si.txt", "attachments/bl.pdf"],
        }
        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
            "attachments/bl.pdf": b"not a real pdf, will not open",
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertEqual(result.review_reason, "unreadable")
        self.assertNotEqual(result.review_reason, "missing_attachment")
        self.assertEqual(result.status, "NEEDS_REVIEW")

    def test_5_valid_si_and_bl_proceeds_normally(self):
        email = {
            "email_id": "email_test_valid_pair",
            "from": "ops@example.com",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARISON_BODY,
            "attachments": ["attachments/si.txt", "attachments/bl.txt"],
        }
        attachment_bytes = {
            "attachments/si.txt": SI_TEXT.encode(),
            "attachments/bl.txt": BL_TEXT.encode(),
        }

        result = process_email(email, attachment_bytes=attachment_bytes)

        self.assertIsNone(result.review_reason)
        self.assertEqual(result.status, "OK")
        self.assertIsNotNone(result.si)
        self.assertIsNotNone(result.bl)

    def test_6_no_attachments_no_si_reference_resolves_ok_not_review(self):
        # The core regression: a plain "please send the draft BL" request
        # (no attachments, no mention of an SI to compare against) must
        # resolve cleanly, not be escalated as a missing attachment.
        email = {
            "email_id": "email_test_send_bl_only",
            "from": "ops@example.com",
            "subject": SEND_DRAFT_BL_SUBJECT,
            "body": SEND_DRAFT_BL_BODY,
            "attachments": [],
        }

        with patch("pipeline.classify.classify_with_llm", return_value="BL_COMPARISON"):
            result = process_email(email, attachment_bytes={})

        self.assertEqual(result.category, "BL_COMPARISON")
        self.assertEqual(result.status, "OK")
        self.assertIsNone(result.review_reason)
        self.assertFalse(result.has_defect)
        self.assertEqual(result.defect_fields, [])

    def test_7_no_attachments_with_si_reference_still_needs_review(self):
        # Genuine missing-attachment behavior must be preserved: a
        # comparison explicitly requested with nothing attached still
        # escalates, even with zero physical attachments.
        email = {
            "email_id": "email_test_dropped_attachments",
            "from": "ops@example.com",
            "subject": COMPARISON_SUBJECT,
            "body": COMPARE_NO_ATTACHMENT_BODY,
            "attachments": [],
        }

        result = process_email(email, attachment_bytes={})

        self.assertEqual(result.status, "NEEDS_REVIEW")
        self.assertEqual(result.review_reason, "missing_attachment")
        self.assertFalse(result.has_defect)
        self.assertEqual(result.defect_fields, [])


if __name__ == "__main__":
    unittest.main()