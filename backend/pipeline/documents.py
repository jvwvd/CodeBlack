from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook

from llm import read_scanned_pdf_with_vision


_REPO_ROOT = Path(__file__).resolve().parents[2]


def _resolve_path(path: str) -> Path | None:
    supplied = Path(path)

    candidates = [
        supplied,
        Path.cwd() / supplied,
        _REPO_ROOT / supplied,
        _REPO_ROOT / "data" / supplied,
    ]

    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate

    return None


def _parse_txt(data: bytes) -> str | None:
    try:
        return data.decode("utf-8", errors="replace")
    except Exception:
        return None


def _parse_pdf(data: bytes) -> str | None:
    """
    Native PDF extraction first.
    Vision fallback only when native text is insufficient.
    """

    try:
        document = pymupdf.open(
            stream=data,
            filetype="pdf",
        )

        text = "\n".join(
            page.get_text("text")
            for page in document
        )

        document.close()

    except Exception:
        return None

    native_text = text.strip()

    # Native PDF: cheapest and deterministic path.
    if len(native_text) >= 20:
        return native_text

    # Image-only/scanned PDF: Gemini vision fallback.
    vision_text = read_scanned_pdf_with_vision(data)

    if vision_text and len(vision_text.strip()) >= 20:
        return vision_text.strip()

    # Existing reliability logic will route this to NEEDS_REVIEW.
    return None


def _parse_docx(data: bytes) -> str | None:
    try:
        document = Document(BytesIO(data))
        chunks: list[str] = []

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if value:
                chunks.append(value)

        for table in document.tables:
            for row in table.rows:
                values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if values:
                    chunks.append(" | ".join(values))

        text = "\n".join(chunks)
        return text.strip() or None

    except Exception:
        return None


def _parse_xlsx(data: bytes) -> str | None:
    try:
        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=True,
        )

        chunks: list[str] = []

        for worksheet in workbook.worksheets:
            chunks.append(f"[Sheet: {worksheet.title}]")

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None and str(value).strip()
                ]

                if values:
                    chunks.append(" | ".join(values))

        workbook.close()

        text = "\n".join(chunks)
        return text.strip() or None

    except Exception:
        return None


def parse_attachment_bytes(
    data: bytes,
    filename: str,
) -> str | None:
    """
    Parse raw attachment bytes.

    Used by both local files and cloud/Supabase integration.
    """

    extension = Path(filename).suffix.lower()

    if extension == ".txt":
        return _parse_txt(data)

    if extension == ".pdf":
        return _parse_pdf(data)

    if extension == ".docx":
        return _parse_docx(data)

    if extension == ".xlsx":
        return _parse_xlsx(data)

    return None


def read_attachment(path: str) -> str | None:
    """
    Frozen local/evaluation contract.

    Returns extracted text or None if unsupported/unreadable.
    """

    resolved = _resolve_path(path)

    if resolved is None:
        return None

    try:
        data = resolved.read_bytes()
    except OSError:
        return None

    return parse_attachment_bytes(data, resolved.name)


def identify_doc_type(text: str) -> str:
    if not text or len(text.strip()) < 20:
        return "UNREADABLE"

    upper = text.upper()

    # SI documents may be titled Shipping Instruction,
    # BL Instruction, or Bill of Lading Instruction.
    if (
        "SHIPPING INSTRUCTION" in upper
        or "SHIPPING INSTRUCTIONS" in upper
        or "BILL OF LADING INSTRUCTION" in upper
        or "B/L INSTRUCTION" in upper
        or "BL INSTRUCTION" in upper
    ):
        return "SI"

    # Draft BL documents.
    if (
        "BILL OF LADING (DRAFT)" in upper
        or "DRAFT BILL OF LADING" in upper
        or "DRAFT B/L" in upper
        or "DRAFT BL" in upper
    ):
        return "BL"

    # Generic BL, but never mistake an instruction for the BL itself.
    if "BILL OF LADING" in upper and "INSTRUCTION" not in upper:
        return "BL"

    return "OTHER"


def load_pair(
    email: dict,
    attachment_bytes: dict[str, bytes] | None = None,
) -> tuple[str | None, str | None]:
    """
    Find one SI and one BL.

    attachment_bytes allows the deployed backend to supply raw bytes
    downloaded from private Supabase Storage.

    Local evaluation continues to use read_attachment().
    """

    si_text: str | None = None
    bl_text: str | None = None
    unreadable_found = False

    for attachment in email.get("attachments", []):
        text: str | None

        if attachment_bytes is not None and attachment in attachment_bytes:
            text = parse_attachment_bytes(
                attachment_bytes[attachment],
                attachment,
            )
        else:
            text = read_attachment(attachment)

        if text is None:
            unreadable_found = True
            continue

        doc_type = identify_doc_type(text)

        if doc_type == "SI" and si_text is None:
            si_text = text

        elif doc_type == "BL" and bl_text is None:
            bl_text = text

    if unreadable_found:
        if si_text is None:
            si_text = ""

        if bl_text is None:
            bl_text = ""

    return si_text, bl_text